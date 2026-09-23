from __future__ import annotations

import os
import re
import tempfile
import math
import unicodedata
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill

from crm_report_exporter.exporter import (
    EXPECTED_TEMPLATE_SHEET,
    _apply_report_dates,
    _clear_formulas_and_invalid_names,
    _prepare_new_sheet,
    _scan_error_tokens,
    _scan_missing_report_values,
    _style_body_row,
    _validate_metric_bundle,
    _write_brand_metrics,
    _write_district_metrics,
    _write_mall_and_capture,
    _write_member_sections,
)
from crm_report_exporter.models import ReportExportError, ReportExportResult

FIELD_LABELS = {
    "description": "说明",
    "evidence_strength": "证据强度",
    "evidence_ids": "证据编号",
}
EVIDENCE_STRENGTH_LABELS = {
    "DIRECT": "直接证据",
    "AGGREGATE": "聚合证据",
    "INSUFFICIENT": "证据不足",
}
MANDATORY_REASON_LABELS = {
    "CORE_SCOPE_RED_OR_YELLOW": "核心范围红/黄异常",
    "MATERIAL_ISSUE": "经营影响筛选",
    "PRIORITY_BRAND_RED_OR_YELLOW": "重点品牌红/黄异常",
    "SEVERE_DATA_QUALITY": "严重数据质量问题",
}
DIMENSION_LABELS = {"TOTAL": "总体", "REGION": "街区", "MEMBER_TIER": "卡级", "REGION_MEMBER_TIER": "街区×卡级", "BRAND": "品牌"}
METRIC_LABELS = {
    "CRM_SALES": "CRM 销售额",
    "MALL_SALES": "商场销售额",
    "CAPTURE_RATIO": "Capture Ratio",
    "ACTIVE_MEMBERS": "积分消费会员数",
    "NEW_REGISTRATIONS": "新增会员数",
    "TRANSACTIONS": "交易笔数",
    "AVG_TICKET": "笔单价",
    "CARD_TIER_CRM_SALES": "品牌卡级CRM销售",
    "AMOUNT_BAND_CRM_SALES": "高奢金额段CRM销售",
    "REFUND_ADJUSTMENT": "净额调整项",
    "WEEKLY_HISTORY": "自然周销售历史",
    "SALES_PER_ACTIVE_MEMBER": "会员人均消费",
    "BRAND_CARD_TIER_ACTIVE_MEMBERS": "品牌卡级消费会员数",
    "NEW_EXISTING_ACTIVE_MEMBERS": "新老客消费会员数",
    "LARGE_SPEND_CONTRIBUTION": "大额消费贡献",
    "ACTIVITY_RESPONSE": "活动响应",
}
CONFIDENCE_LABELS = {"HIGH": "高", "MEDIUM": "中", "LOW": "低"}
IMPACT_LABELS = {
    "REAL_OPERATION": "真实经营变化",
    "POINTS_PARTICIPATION": "会员积分参与变化",
    "MIXED": "真实经营与会员积分参与共同影响",
    "UNRESOLVED": "暂无法区分真实经营与会员积分参与影响",
}
COMPARISON_LABELS = {
    "WOW": "环比",
    "YOY": "同比",
    "YTD": "累计同比",
    "CONTINUOUS_PERIOD": "连续期",
    "EIGHT_WEEK_BASELINE": "八周基线",
}
REGION_LABELS = {"N": "北区", "S": "南区", "W": "西区"}


def _mapping(value: Any, path: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise ReportExportError("FIELD_INVALID", f"{path}必须是JSON对象")
    return value


def _list(value: Any, path: str) -> list[Any]:
    if not isinstance(value, list):
        raise ReportExportError("FIELD_INVALID", f"{path}必须是数组")
    return value


def _join(value: Any) -> str:
    if isinstance(value, list):
        return "；".join(_join(item) for item in value)
    if isinstance(value, Mapping):
        parts = []
        for key, item in value.items():
            if item in (None, "", []):
                continue
            rendered = _join(item)
            if key == "evidence_strength":
                rendered = EVIDENCE_STRENGTH_LABELS.get(str(item).upper(), rendered)
            parts.append(f"{FIELD_LABELS.get(str(key), str(key))}：{rendered}")
        return "；".join(parts)
    return str(value or "")


def _narrative_text(value: Any) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return "；".join(part for part in (_narrative_text(item) for item in value) if part)
    if isinstance(value, Mapping):
        for key in ("statement", "description", "conclusion", "reason", "finding"):
            text = str(value.get(key, "")).strip()
            if text:
                return text
        return "；".join(
            part
            for key, item in value.items()
            if key not in {"evidence_ids", "evidence_strength", "applicability_limits"}
            for part in [_narrative_text(item)]
            if part
        )
    return str(value or "").strip()


def _sentence(value: Any) -> str:
    text = _narrative_text(value).rstrip("。；;，, ")
    return f"{text}。" if text else ""


def _object_name(group: Mapping[str, Any]) -> str:
    dimension = str(group.get("dimension_type", "")).upper()
    key = str(group.get("dimension_key", "")).strip()
    if dimension == "TOTAL" or key.upper() == "TOTAL":
        return "总盘"
    if dimension == "REGION":
        return REGION_LABELS.get(key.upper(), key)
    if dimension in {"BRAND", "STORE"}:
        return f"品牌 {key}"
    if dimension in {"MEMBER_TIER", "CARD", "CARD_TIER"}:
        return key if key.endswith("卡") else f"{key}卡级"
    if dimension in {"REGION_MEMBER_TIER", "REGION_CARD", "REGION_CARD_TIER"}:
        parts = [part.strip() for part in key.replace("/", "|").split("|", 1)]
        if len(parts) == 2:
            region = REGION_LABELS.get(parts[0].upper(), parts[0])
            tier = parts[1] if parts[1].endswith("卡") else f"{parts[1]}卡级"
            return f"{region}的{tier}"
    label = DIMENSION_LABELS.get(dimension, dimension or "对象")
    return f"{label} {key}".strip()


def _anomaly_object_metrics(insight: Mapping[str, Any], groups: list[Mapping[str, Any]]) -> str:
    group_index = {str(group.get("anomaly_group_id")): group for group in groups}
    values: list[str] = []
    for group_id in insight.get("anomaly_group_ids", []):
        group = group_index.get(str(group_id))
        if not group:
            continue
        metric = METRIC_LABELS.get(str(group.get("metric_id", "")).upper(), str(group.get("metric_id", "")).strip() or "指标")
        value = f"{_object_name(group)}的{metric}"
        if value not in values:
            values.append(value)
    return "\n".join(values) or "异常对象和指标待补充"


def _period_text(periods: Any) -> str:
    if not isinstance(periods, Mapping):
        return ""
    current = periods.get("current")
    if isinstance(current, list) and current:
        if len(current) == 1:
            return str(current[0])
        return f"{current[0]}至{current[-1]}"
    return ""


def _direction_text(record: Mapping[str, Any]) -> str:
    change = record.get("change")
    if isinstance(change, (int, float)) and not isinstance(change, bool):
        return "上升" if change > 0 else "下降" if change < 0 else "持平"
    display = str(record.get("display_change", "")).strip()
    return "上升" if display.startswith("+") else "下降" if display.startswith("-") else ""


def _scope_name(scope: Any) -> str:
    text = str(scope or "").strip()
    if not text or text.upper() == "TOTAL":
        return "总盘"
    if ":" not in text:
        return text
    dimension, key = text.split(":", 1)
    return _object_name({"dimension_type": dimension, "dimension_key": key})


def _fact_description(fact: Mapping[str, Any]) -> str:
    scope = _scope_name(fact.get("scope"))
    metric_id = str(fact.get("metric_id", fact.get("metric", ""))).upper()
    metric = METRIC_LABELS.get(metric_id, metric_id or "指标")
    details = [
        COMPARISON_LABELS.get(str(fact.get("comparison_basis", "")).upper(), str(fact.get("comparison_basis", "")).strip()),
        _period_text(fact.get("periods")),
        _direction_text(fact),
    ]
    qualifier = "、".join(item for item in details if item)
    return f"{scope}的{metric}（{qualifier}）" if qualifier else f"{scope}的{metric}"


def _activity_description(activity: Mapping[str, Any]) -> str:
    name = str(activity.get("activity_name", activity.get("activity_id", "活动"))).strip()
    start = str(activity.get("start_date", "")).strip()
    end = str(activity.get("end_date", "")).strip()
    date_text = f"{start}至{end}" if start and end and start != end else start or end
    reasons = activity.get("match_reasons", [])
    reason_text = "、".join(str(item).strip() for item in reasons if str(item).strip()) if isinstance(reasons, list) else ""
    details = "；".join(item for item in (date_text, reason_text) if item)
    return f"{name}（{details}）" if details else name


def _output_insight_text(insight: Mapping[str, Any]) -> str:
    text = insight.get("narrative")
    if not isinstance(text, str) or not text.strip():
        raise ReportExportError("NARRATIVE_REQUIRED", "需要模型撰写的完整业务正文")
    action = insight.get("action_text")
    if not isinstance(action, str) or not action.strip() or not text.strip().endswith("\n\n" + action.strip()):
        raise ReportExportError("ACTION_NOT_DELIVERED", "完整行动段必须进入业务正文；请通过新版写作入口组装")
    if insight.get("recommendation", {}).get("applicable") is True and insight["recommendation"].get("action", "").strip() != action.strip():
        raise ReportExportError("ACTION_TEXT_MISMATCH", "业务正文与内部行动不一致")
    return text.strip()


def _validate_v3(
    metric_bundle: Mapping[str, Any],
    anomaly_bundle: Mapping[str, Any],
    plan: Mapping[str, Any],
    insight_bundle: Mapping[str, Any],
    evidence_bundle: Mapping[str, Any],
) -> tuple[list[Mapping[str, Any]], list[Mapping[str, Any]], list[Mapping[str, Any]]]:
    artifacts = {
        "AnomalyBundle": anomaly_bundle,
        "InvestigationPlan": plan,
        "EvidenceLedger": evidence_bundle,
        "InsightBundle": insight_bundle,
    }
    expected = {field: metric_bundle.get(field) for field in ("week_id", "input_signature", "implementation_version")}
    for name, artifact in artifacts.items():
        if str(artifact.get("schema_version")) != "3.0":
            raise ReportExportError("SCHEMA_VERSION_UNSUPPORTED", f"{name}必须为3.0")
        mismatches = {
            field: {"expected": expected[field], "actual": artifact.get(field)}
            for field in expected
            if not str(artifact.get(field, "")).strip() or artifact.get(field) != expected[field]
        }
        if mismatches:
            raise ReportExportError("ARTIFACT_IDENTITY_MISMATCH", f"{name}与MetricBundle身份不一致", mismatches)
    if insight_bundle.get("kb_version") != evidence_bundle.get("kb_version"):
        raise ReportExportError("ARTIFACT_IDENTITY_MISMATCH", "InsightBundle与EvidenceLedger知识库版本不一致")
    groups = [_mapping(item, "anomaly_group") for item in _list(anomaly_bundle.get("anomaly_groups"), "anomaly_groups")]
    group_id_list = [str(item.get("anomaly_group_id", "")).strip() for item in groups]
    if any(not item for item in group_id_list) or len(group_id_list) != len(set(group_id_list)):
        raise ReportExportError("ANOMALY_REFERENCE_INVALID", "异常组ID必须非空且唯一")
    group_ids = set(group_id_list)
    selected_cases = _list(plan.get("selected_cases"), "selected_cases")
    selected_ids = [str(case.get("insight_case_id", "")).strip() for case in selected_cases]
    if any(not item for item in selected_ids) or len(selected_ids) != len(set(selected_ids)):
        raise ReportExportError("CASE_REFERENCE_INVALID", "selected case ID必须非空且唯一")
    selected = {str(case.get("insight_case_id")): {str(item) for item in case.get("anomaly_group_ids", [])} for case in selected_cases}
    selected_groups = set().union(*selected.values(), set())
    omitted_items = _list(plan.get("omitted_anomaly_groups"), "omitted_anomaly_groups")
    omitted_ids = [str(item.get("anomaly_group_id", "")).strip() for item in omitted_items]
    if any(not item for item in omitted_ids) or len(omitted_ids) != len(set(omitted_ids)) or selected_groups & set(omitted_ids) or selected_groups | set(omitted_ids) != group_ids:
        raise ReportExportError("CANDIDATE_ACCOUNTING_INVALID", "完整候选必须且只能入选或省略一次")
    mandatory_ids = {str(item.get("anomaly_group_id")) for item in groups if item.get("must_investigate") is True}
    if not mandatory_ids.issubset(selected_groups):
        raise ReportExportError("MANDATORY_OMITTED", "业务必查异常不得省略")
    insights = [_mapping(item, "insight") for item in _list(insight_bundle.get("insights"), "insights")]
    monitors = [_mapping(item, "monitor_item") for item in _list(insight_bundle.get("monitor_items"), "monitor_items")]
    reporting_policy = metric_bundle.get("reporting_policy", {})
    if not isinstance(reporting_policy, Mapping):
        raise ReportExportError("REPORTING_POLICY_INVALID", "reporting_policy必须是对象")
    if len(insight_bundle.get("summary_items",[]))>5:
        raise ReportExportError("SUMMARY_LIMIT","摘要最多五点")
    if len(insights)>8 and not insight_bundle.get("deduplication_review",{}).get("completed"):
        raise ReportExportError("INSIGHT_REVIEW_REQUIRED","超过八条需要去重复核")
    decision_ids = [str(item.get("insight_case_id", "")).strip() for item in insights + monitors]
    decisions = set(decision_ids)
    if any(not item for item in decision_ids) or len(decision_ids) != len(decisions) or decisions != set(selected):
        raise ReportExportError("CASE_COUNT_MISMATCH", "所有selected case必须在报告或监控中出现一次")
    for case_id, members in selected.items():
        if not members or not members.issubset(group_ids):
            raise ReportExportError("CASE_REFERENCE_INVALID", "case引用了不存在的异常", {"insight_case_id": case_id})
    for record in insights + monitors:
        case_id = str(record.get("insight_case_id"))
        if {str(item) for item in record.get("anomaly_group_ids", [])} != selected[case_id]:
            raise ReportExportError("CASE_REFERENCE_INVALID", "最终决策不得改变case成员", {"insight_case_id": case_id})
    monitor_case_ids = {str(item.get("insight_case_id")) for item in monitors}
    total_red_ids = {
        str(item.get("anomaly_group_id"))
        for item in groups
        if str(item.get("dimension_type", "")).upper() == "TOTAL"
        and str(item.get("severity", item.get("group_severity", ""))).upper() == "RED"
    }
    for case_id in monitor_case_ids:
        if selected[case_id] & total_red_ids and not insight_bundle.get("data_quality_items"):
            raise ReportExportError("TOTAL_RED_MONITOR_ONLY", "含Total红色异常的case必须优先REPORT", {"insight_case_id": case_id})
    evidence_cases = [_mapping(item, "evidence_case") for item in _list(evidence_bundle.get("cases"), "evidence.cases")]
    evidence_case_ids = [str(item.get("insight_case_id", "")).strip() for item in evidence_cases]
    if len(evidence_case_ids) != len(set(evidence_case_ids)) or set(evidence_case_ids) != set(selected):
        raise ReportExportError("COVERAGE_MISMATCH", "EvidenceLedger必须且只能覆盖全部selected case")
    valid_statuses = {"INVESTIGATED", "UPSTREAM_SUFFICIENT", "DATA_UNAVAILABLE", "NOT_APPLICABLE"}
    coverage_ids: list[str] = []
    for case in evidence_cases:
        case_id = str(case.get("insight_case_id"))
        if {str(item) for item in case.get("anomaly_group_ids", [])} != selected[case_id]:
            raise ReportExportError("CASE_REFERENCE_INVALID", "EvidenceLedger不得改变case成员", {"insight_case_id": case_id})
        for coverage in case.get("coverage", []):
            group_id = str(coverage.get("anomaly_group_id", ""))
            if group_id not in selected[case_id] or str(coverage.get("status", "")) not in valid_statuses:
                raise ReportExportError("COVERAGE_MISMATCH", "调查覆盖状态或归属无效", {"insight_case_id": case_id, "anomaly_group_id": group_id})
            coverage_ids.append(group_id)
    if set(coverage_ids) != selected_groups:
        raise ReportExportError("COVERAGE_MISMATCH", "每个入选异常必须且只能有一个最终覆盖状态")
    return groups, insights, monitors


def _write_insights(sheet, report, groups, insights, evidence):
    _prepare_new_sheet(sheet, report, ["主题", "洞察"], [28, 112])
    for row_index, insight in enumerate(insights, 2):
        _style_body_row(sheet, row_index, 2, report)
        sheet.cell(row_index, 1, insight["title"])
        sheet.cell(row_index, 2, _output_insight_text(insight))
    if not insights:
        sheet.cell(2, 1, "本周无需要单独展开的重要变化")
    sheet.auto_filter.ref = f"A1:B{max(2, sheet.max_row)}"
    sheet.print_area = sheet.auto_filter.ref
    sheet.freeze_panes = "A2"
    return len(insights)


def _write_analysis_notes(sheet, report, insights, evidence, insight_bundle, metric_bundle):
    _prepare_new_sheet(sheet, report, ["主题", "置信度", "判断依据", "资料来源", "必要说明"], [28, 10, 56, 42, 62])
    facts = {f["fact_id"]: f for c in evidence.get("cases", []) for f in c.get("facts", [])}
    activities = {a["activity_id"]: a for c in evidence.get("cases", []) for a in c.get("activities", [])}
    for insight in insights:
        sources = []
        if insight.get("fact_ids"):
            sources.append("CRM消费记录")
        if any("MALL" in str(facts.get(fid, {}).get("metric_id", "")) for fid in insight.get("fact_ids", [])):
            sources.append("商场销售表")
        for aid in insight.get("activity_ids", []):
            activity = activities.get(aid, {})
            if activity.get("activity_name"):
                sources.append(str(activity["activity_name"]))
        note = insight.get("analysis_note", {})
        sheet.append([insight["title"], CONFIDENCE_LABELS.get(insight["confidence"], insight["confidence"]),
                      note.get("basis", ""), "；".join(dict.fromkeys(sources)), note.get("limitations", "")])
    # Shared limitations once, without raw IDs, missing-cell lists or audit records.
    shared = insight_bundle.get("report_notes", [])
    if not shared:
        shared = list(dict.fromkeys(str(x.get("message", "")) for x in metric_bundle.get("data_quality_status", []) if x.get("message")))
    for note in dict.fromkeys(shared):
        sheet.append(["报告口径", "", "", "源数据与报告口径", note])
    # Material topics disclosed via data quality must remain visible in the slim report.
    for item in insight_bundle.get("data_quality_items", []):
        if item.get("anomaly_group_ids"):
            sheet.append([item.get("affected_judgments", "数据限制"), "", "", "源数据", item["reason"]])
    sheet.freeze_panes = "A2"
    sheet.print_area = f"A1:E{sheet.max_row}"


def _fit_text_rows(sheet: Any) -> None:
    """Keep long business narratives visible within Excel's row-height limit."""
    for row in sheet.iter_rows(min_row=2):
        required_height = 24.0
        for cell in row:
            if cell.value is None:
                continue
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            width = max(8.0, (sheet.column_dimensions[cell.column_letter].width or 13) - 3)
            lines = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1
                                             for ch in part) / width))
                        for part in str(cell.value).split('\n'))
            # Mixed Latin/Chinese paragraphs can wrap at word boundaries before
            # the estimated character width, so reserve two lines for long prose.
            if sheet.title == '洞察' and len(str(cell.value)) > 100:
                lines += 2
            required_height = max(required_height, lines * max(15.0, float(cell.font.sz or 11) * 1.5) + 8)
        sheet.row_dimensions[row[0].row].height = min(409.5, required_height)


def export_weekly_report_v3(
    template_path: str | Path,
    output_path: str | Path,
    metric_bundle: Mapping[str, Any],
    anomaly_bundle: Mapping[str, Any],
    plan: Mapping[str, Any],
    insight_bundle: Mapping[str, Any],
    evidence_bundle: Mapping[str, Any],
) -> ReportExportResult:
    template = Path(template_path).expanduser().resolve()
    output = Path(output_path).expanduser().resolve()
    if not template.is_file() or template == output:
        raise ReportExportError("TEMPLATE_INVALID", "模板必须存在且输出不得覆盖模板", {"template": str(template), "output": str(output)})
    metrics, windows = _validate_metric_bundle(metric_bundle)
    groups, insights, monitors = _validate_v3(metric_bundle, anomaly_bundle, plan, insight_bundle, evidence_bundle)
    workbook = load_workbook(template)
    if workbook.sheetnames != [EXPECTED_TEMPLATE_SHEET]:
        raise ReportExportError("TEMPLATE_STRUCTURE_INVALID", "模板必须且只能包含Weekly Summary")
    report = workbook[EXPECTED_TEMPLATE_SHEET]
    report.title = "报告"
    _clear_formulas_and_invalid_names(workbook, report)
    _apply_report_dates(report, windows)
    _write_district_metrics(report, metrics)
    _write_brand_metrics(report, metrics)
    _write_mall_and_capture(report, metrics)
    _write_member_sections(report, metrics)
    # Normalize display labels without modifying the source template.
    for row in report:
        for cell in row:
            if isinstance(cell.value, str) and cell.data_type != 'f':
                cell.value = re.sub(r'\bcrm\s*sales\b', 'crm sales', cell.value, flags=re.I)
                cell.value = re.sub(r'\bmall\s*sales\b', 'mall sales', cell.value, flags=re.I)
    report.sheet_view.showGridLines = False
    report.print_area = "B2:J109"
    report.page_setup.orientation = "landscape"
    report.page_setup.fitToWidth = 1
    report.page_setup.fitToHeight = 0
    report.sheet_properties.pageSetUpPr.fitToPage = True
    insight_count = _write_insights(workbook.create_sheet("洞察"), report, groups, insights, evidence_bundle)
    alert_count = 0
    _write_analysis_notes(workbook.create_sheet("分析说明"), report, insights, evidence_bundle, insight_bundle, metric_bundle)

    missing = _scan_missing_report_values(report)
    unresolved=[item for item in missing if item["value"].strip().upper()!="N/A"]
    if unresolved:
        raise ReportExportError("REPORT_VALUE_MISSING", "报告仍有未计算指标", {"missing_values": unresolved[:30]})
    for sheet in workbook:
        if sheet.title != "报告":
            _fit_text_rows(sheet)

    errors = _scan_error_tokens(workbook)
    if errors:
        raise ReportExportError("WORKBOOK_ERROR_VALUE", "输出工作簿包含错误值", {"errors": errors[:20]})
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, output)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()
    return ReportExportResult(output_path=output, week_id=str(metric_bundle.get("week_id", "")), sheet_names=tuple(workbook.sheetnames), insight_count=insight_count, alert_count=alert_count, quality_count=len(metric_bundle.get("data_quality_status", [])) + len(anomaly_bundle.get("data_quality_status", [])))
