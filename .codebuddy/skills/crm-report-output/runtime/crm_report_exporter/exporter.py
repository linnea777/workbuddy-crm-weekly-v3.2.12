from __future__ import annotations

import os
import re
import tempfile
from collections.abc import Mapping, Sequence
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .models import ReportExportError, ReportExportResult


EXPECTED_TEMPLATE_SHEET = "Weekly Summary"
METRIC_IDS = tuple(f"M{number:02d}" for number in range(1, 15))
DISTRICT_ROWS = {"S": 3, "N": 4, "W": 5, "Total": 6}
MALL_ROWS = {"S": 21, "N": 22, "W": 23, "Total": 24}
CAPTURE_ROWS = {"S": 27, "N": 28, "W": 29}
CARD_ROWS = {"星钻卡": 94, "黑钻卡": 95, "黑卡": 96, "金卡": 97, "银卡": 98}
NEW_MEMBER_ROWS = {"星钻卡": 40, "黑钻卡": 41, "黑卡": 42, "金卡": 43, "银卡": 44}
POINT_ROWS = {"星钻卡": 32, "黑钻卡": 33, "黑卡": 34, "金卡": 35, "银卡": 36, "Total": 37}
ACTIVE_ROWS = {"星钻卡": 81, "黑钻卡": 82, "黑卡": 83, "金卡": 84, "银卡": 85, "Total": 86}
BRAND_ROWS = {
    "LV": (9, 10),
    "Dior": (13, 14),
    "Hermes": (17, 18),
}
BRAND_FULL_NAMES = {
    "DIOR": "Christian Dior（迪奥）",
    "HERMES": "Hermès（爱马仕）",
    "LV": "Louis Vuitton（路易威登）",
}
CONFIDENCE_LABELS = {"HIGH": "高", "MEDIUM": "中", "LOW": "低"}
SEVERITY_LABELS = {"RED": "红色", "YELLOW": "黄色", "GREEN": "绿色"}
DIRECTION_LABELS = {"UP": "上升", "DOWN": "下降", "FLAT": "持平"}
DIMENSION_LABELS = {
    "TOTAL": "总体",
    "REGION": "街区",
    "DISTRICT": "街区",
    "BRAND": "品牌",
    "STORE": "店铺",
    "CARD_LEVEL": "卡级",
    "REGION_CARD_LEVEL": "街区×卡级",
}
METRIC_LABELS = {
    "CRM_SALES": "CRM销售额",
    "MALL_SALES": "Mall Sales",
    "CAPTURE_RATIO": "Capture Ratio",
    "MEMBER_COUNT": "会员人数",
    "TRANSACTIONS": "交易笔数",
    "AVERAGE_TICKET": "笔单价",
}
SIGNAL_LABELS = {
    "WOW_THRESHOLD": "环比变化",
    "YOY_THRESHOLD": "同比变化",
    "YTD_THRESHOLD": "年累计同比变化",
    "CONTINUOUS_DECLINE": "连续下降",
    "EIGHT_WEEK_MEDIAN_DEVIATION": "偏离近8周常态",
    "ROBUST_ZSCORE": "超出历史正常波动",
    "CAPTURE_RATIO_PP_CHANGE": "Capture Ratio变化",
}
MEMBER_TYPE_LABELS = {"new": "新会员", "existing": "老会员", "old": "老会员", "新会员": "新会员", "老会员": "老会员"}
NEW_MEMBER_FONT_COLOR = "FF0000FF"
STAR_DIAMOND_FONT_COLOR = "FF7030A0"
DEFAULT_FONT_COLOR = "FF000000"
ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")
MISSING_VALUE_TOKENS = {"N/A", "NA"}


def _require_mapping(value: Any, name: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise ReportExportError("INVALID_BUNDLE", f"{name}必须是Mapping", {"bundle": name})
    return value


def _require_sequence(value: Any, path: str) -> Sequence[Any]:
    if not isinstance(value, Sequence) or isinstance(value, (str, bytes, bytearray)):
        raise ReportExportError("INVALID_BUNDLE", f"{path}必须是数组", {"path": path})
    return value


def _validate_metric_bundle(bundle: Mapping[str, Any]) -> tuple[Mapping[str, Any], Mapping[str, Any]]:
    metrics = _require_mapping(bundle.get("metrics"), "metric_bundle.metrics")
    missing = [metric_id for metric_id in METRIC_IDS if metric_id not in metrics]
    if missing:
        raise ReportExportError("METRIC_MISSING", "MetricBundle缺少必要指标", {"missing_metrics": missing})
    windows = _require_mapping(bundle.get("windows"), "metric_bundle.windows")
    for window_name in ("current", "previous", "yoy", "ytd", "previous_ytd"):
        window = _require_mapping(windows.get(window_name), f"metric_bundle.windows.{window_name}")
        if not window.get("start") or not window.get("end"):
            raise ReportExportError("WINDOW_MISSING", "比较周期缺少起止日期", {"window": window_name})
    return metrics, windows


def _validate_anomaly_bundle(bundle: Mapping[str, Any]) -> tuple[list[Mapping[str, Any]], list[Mapping[str, Any]]]:
    version = bundle.get("schema_version")
    if version is not None and version != "1.1":
        raise ReportExportError("SCHEMA_VERSION_UNSUPPORTED", "仅支持AnomalyBundle 1.1", {"actual": version})
    groups = [_require_mapping(item, "anomaly_group") for item in _require_sequence(bundle.get("anomaly_groups", []), "anomaly_groups")]
    signals = [_require_mapping(item, "anomaly_signal") for item in _require_sequence(bundle.get("anomaly_signals", []), "anomaly_signals")]
    return groups, signals


def _text_list(value: Any, path: str, *, allow_empty: bool = True) -> list[str]:
    values = _require_sequence(value, path)
    result = []
    for index, item in enumerate(values):
        if not isinstance(item, str) or not item.strip():
            raise ReportExportError(
                "INSIGHT_FIELD_INVALID",
                f"{path}只能包含非空字符串",
                {"path": f"{path}[{index}]"},
            )
        result.append(item.strip())
    if not allow_empty and not result:
        raise ReportExportError("INSIGHT_FIELD_INVALID", f"{path}不能为空", {"path": path})
    if len(result) != len(set(result)):
        raise ReportExportError("INSIGHT_FIELD_INVALID", f"{path}不能包含重复值", {"path": path})
    return result


def _validate_insights(
    bundle: Mapping[str, Any] | None,
    groups: list[Mapping[str, Any]],
    signals: list[Mapping[str, Any]],
) -> list[Mapping[str, Any]] | None:
    if bundle is None:
        return None
    version = bundle.get("schema_version")
    if version != "2.0":
        raise ReportExportError("SCHEMA_VERSION_UNSUPPORTED", "仅支持InsightBundle 2.0", {"actual": version})
    knowledge_base_status = bundle.get("knowledge_base_status")
    if knowledge_base_status not in ("available", "unavailable"):
        raise ReportExportError(
            "INSIGHT_FIELD_INVALID",
            "knowledge_base_status必须是available/unavailable",
            {"actual": knowledge_base_status},
        )
    group_ids = {str(group.get("anomaly_group_id", "")) for group in groups}
    signal_owners = {
        str(signal.get("signal_id", "")): str(signal.get("anomaly_group_id", ""))
        for signal in signals
    }
    insights = [_require_mapping(item, "insight") for item in _require_sequence(bundle.get("insights"), "insights")]
    seen_group_ids: set[str] = set()
    for index, insight in enumerate(insights):
        path = f"insights[{index}]"
        missing = [
            field
            for field in (
                "anomaly_group_id",
                "signal_ids",
                "data_fact",
                "fact_ids",
                "evidence_refs",
                "inference",
                "confidence",
                "confidence_reason",
                "recommendation",
                "verification_metrics",
            )
            if field not in insight
        ]
        if missing:
            raise ReportExportError(
                "INSIGHT_FIELD_MISSING",
                "InsightRecord缺少必要字段",
                {"index": index, "missing_fields": missing},
            )
        group_id = str(insight["anomaly_group_id"]).strip()
        if not group_id or group_id not in group_ids:
            raise ReportExportError(
                "INSIGHT_REFERENCE_INVALID",
                "洞察引用了不存在的anomaly_group_id",
                {"index": index, "anomaly_group_id": group_id},
            )
        if group_id in seen_group_ids:
            raise ReportExportError(
                "INSIGHT_REFERENCE_INVALID",
                "同一异常组只能生成一条洞察",
                {"index": index, "anomaly_group_id": group_id},
            )
        seen_group_ids.add(group_id)

        signal_ids = _text_list(insight["signal_ids"], f"{path}.signal_ids", allow_empty=False)
        invalid_signals = [signal_id for signal_id in signal_ids if signal_owners.get(signal_id) != group_id]
        if invalid_signals:
            raise ReportExportError(
                "INSIGHT_REFERENCE_INVALID",
                "signal_ids必须属于当前异常组",
                {"index": index, "invalid_signal_ids": invalid_signals},
            )
        _require_mapping(insight["data_fact"], f"{path}.data_fact")
        _text_list(insight["fact_ids"], f"{path}.fact_ids")
        evidence_refs = _text_list(insight["evidence_refs"], f"{path}.evidence_refs")
        if knowledge_base_status == "unavailable" and evidence_refs:
            raise ReportExportError(
                "INSIGHT_REFERENCE_INVALID",
                "知识库不可用时evidence_refs必须为空",
                {"index": index},
            )
        _text_list(insight["verification_metrics"], f"{path}.verification_metrics", allow_empty=False)
        if not str(insight["inference"]).strip():
            raise ReportExportError(
                "INSIGHT_FIELD_MISSING",
                "inference不能为空",
                {"index": index, "field": "inference"},
            )
        if not str(insight["confidence_reason"]).strip():
            raise ReportExportError(
                "INSIGHT_FIELD_MISSING",
                "confidence_reason不能为空",
                {"index": index, "field": "confidence_reason"},
            )
        confidence = str(insight["confidence"]).upper()
        if confidence not in CONFIDENCE_LABELS:
            raise ReportExportError("INSIGHT_CONFIDENCE_INVALID", "置信度必须是HIGH/MEDIUM/LOW", {"index": index})
        recommendation = _require_mapping(insight["recommendation"], f"{path}.recommendation")
        recommendation_missing = [
            field
            for field in ("applicable", "target_members", "action", "channel", "timing", "not_applicable_reason")
            if field not in recommendation
        ]
        if recommendation_missing:
            raise ReportExportError(
                "INSIGHT_FIELD_MISSING",
                "recommendation缺少必要字段",
                {"index": index, "missing_fields": recommendation_missing},
            )
        if not isinstance(recommendation["applicable"], bool):
            raise ReportExportError(
                "INSIGHT_FIELD_INVALID",
                "recommendation.applicable必须是布尔值",
                {"index": index},
            )
        if recommendation["applicable"] is True and not str(recommendation["action"]).strip():
            raise ReportExportError(
                "INSIGHT_FIELD_MISSING",
                "建议适用时action不能为空",
                {"index": index},
            )
        if recommendation["applicable"] is False and not str(recommendation["not_applicable_reason"]).strip():
            raise ReportExportError(
                "INSIGHT_FIELD_MISSING",
                "建议不适用时not_applicable_reason不能为空",
                {"index": index},
            )
    return insights


def _index(records: Any, key: str, path: str) -> dict[str, Mapping[str, Any]]:
    result: dict[str, Mapping[str, Any]] = {}
    for record in _require_sequence(records, path):
        item = _require_mapping(record, path)
        if key not in item:
            raise ReportExportError("METRIC_FIELD_MISSING", f"{path}缺少{key}", {"record": dict(item)})
        result[str(item[key])] = item
    return result


def _window_label(window: Mapping[str, Any]) -> str:
    start = date.fromisoformat(str(window["start"]))
    end = date.fromisoformat(str(window["end"]))
    return f"{start.month}.{start.day}-{end.month}.{end.day}"


def _write(cell: Any, value: Any, kind: str = "text") -> None:
    if value is None:
        cell.value = "N/A"
        cell.number_format = "@"
        return
    if kind == "amount" or kind == "count":
        cell.value = float(value) if kind == "amount" else int(value)
        cell.number_format = "#,##0"
    elif kind == "rate_percent":
        cell.value = float(value) / 100
        cell.number_format = "0.0%;[Red]-0.0%"
    elif kind == "ratio":
        cell.value = float(value)
        cell.number_format = "0.0%;[Red]-0.0%"
    elif kind == "pp":
        cell.value = float(value)
        cell.number_format = '0.0 "个百分点";[Red]-0.0 "个百分点"'
    else:
        cell.value = str(value)


def _write_change_rate(cell: Any, record: Mapping[str, Any]) -> None:
    """Render a zero comparison base explicitly instead of as an uncomputed value."""
    rate = record.get("change_rate_percent")
    if rate is None and record.get("current") is not None and record.get("base") == 0:
        cell.value = "基期为0"
        cell.number_format = "@"
        return
    _write(cell, rate, "rate_percent")


def _clear_formulas_and_invalid_names(workbook: Any, report: Worksheet) -> None:
    for row in report.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if cell.data_type == "f" or (isinstance(cell.value, str) and any(token in cell.value for token in ERROR_TOKENS)):
                cell.value = None
    for name in list(workbook.defined_names):
        del workbook.defined_names[name]


def _write_district_metrics(report: Worksheet, metrics: Mapping[str, Any]) -> None:
    m01 = _index(metrics["M01"], "dimension", "metrics.M01")
    m02 = _index(metrics["M02"], "dimension", "metrics.M02")
    m03 = _index(metrics["M03"], "dimension", "metrics.M03")
    for dimension, row in DISTRICT_ROWS.items():
        _write(report[f"C{row}"], m01[dimension].get("current"), "amount")
        _write(report[f"D{row}"], m01[dimension].get("base"), "amount")
        _write(report[f"E{row}"], m01[dimension].get("change_rate_percent"), "rate_percent")
        _write(report[f"F{row}"], m02[dimension].get("base"), "amount")
        _write(report[f"G{row}"], m02[dimension].get("change_rate_percent"), "rate_percent")
        _write(report[f"H{row}"], m03[dimension].get("current"), "amount")
        _write(report[f"I{row}"], m03[dimension].get("base"), "amount")
        _write(report[f"J{row}"], m03[dimension].get("change_rate_percent"), "rate_percent")


def _brand_metrics(metrics: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    result = {"Hermes": _require_mapping(metrics["M10"], "metrics.M10")}
    for item in _require_sequence(metrics["M11"], "metrics.M11"):
        brand = _require_mapping(item, "metrics.M11")
        result[str(brand.get("brand"))] = brand
    return result


def _write_brand_metrics(report: Worksheet, metrics: Mapping[str, Any]) -> None:
    brands = _brand_metrics(metrics)
    for brand_name, (current_row, previous_row) in BRAND_ROWS.items():
        if brand_name not in brands:
            raise ReportExportError("METRIC_FIELD_MISSING", "缺少品牌专题", {"brand": brand_name})
        brand = brands[brand_name]
        crm = _require_mapping(brand.get("crm_sales"), f"{brand_name}.crm_sales")
        transactions = _require_mapping(brand.get("transactions"), f"{brand_name}.transactions")
        ticket = _require_mapping(brand.get("average_ticket"), f"{brand_name}.average_ticket")
        mall = _require_mapping(brand.get("mall_sales"), f"{brand_name}.mall_sales")
        capture = _require_mapping(brand.get("capture_ratio"), f"{brand_name}.capture_ratio")
        _write(report[f"C{current_row}"], crm.get("current"), "amount")
        _write(report[f"D{current_row}"], crm.get("base"), "amount")
        _write(report[f"E{current_row}"], crm.get("change_rate_percent"), "rate_percent")
        _write(report[f"F{current_row}"], transactions.get("current"), "count")
        _write(report[f"G{current_row}"], ticket.get("current"), "amount")
        _write(report[f"H{current_row}"], mall.get("current"), "amount")
        _write(report[f"I{current_row}"], capture.get("current"), "ratio")
        _write(report[f"F{previous_row}"], transactions.get("base"), "count")
        _write(report[f"G{previous_row}"], ticket.get("base"), "amount")
        _write(report[f"H{previous_row}"], mall.get("base"), "amount")
        _write(report[f"I{previous_row}"], capture.get("base"), "ratio")
    _write_brand_segments(report, brands)


def _write_brand_segments(report: Worksheet, brands: Mapping[str, Mapping[str, Any]]) -> None:
    layouts = {
        "LV": (range(49, 55), ("J", "K", "L", "M"), 55),
        "Dior": (range(49, 55), ("Q", "R", "S", "T"), 55),
        "Hermes": (range(64, 70), ("Q", "R", "S", "T"), 70),
    }
    for brand_name, (rows, columns, total_row) in layouts.items():
        segments = [
            _require_mapping(item, f"{brand_name}.segments")
            for item in _require_sequence(brands[brand_name].get("segments", []), f"{brand_name}.segments")
        ]
        display_segments = list(reversed(segments))
        for offset, row in enumerate(rows):
            item = display_segments[offset] if offset < len(display_segments) else {}
            _write(report[f"{columns[0]}{row}"], item.get("members"), "count")
            _write(report[f"{columns[1]}{row}"], item.get("member_share"), "ratio")
            _write(report[f"{columns[2]}{row}"], item.get("sales"), "amount")
            _write(report[f"{columns[3]}{row}"], item.get("sales_share"), "ratio")
        total_members = sum(int(item.get("members") or 0) for item in segments)
        total_sales = sum(float(item.get("sales") or 0) for item in segments)
        _write(report[f"{columns[0]}{total_row}"], total_members, "count")
        _write(report[f"{columns[1]}{total_row}"], 1.0 if total_members else 0.0, "ratio")
        _write(report[f"{columns[2]}{total_row}"], total_sales, "amount")
        _write(report[f"{columns[3]}{total_row}"], 1.0 if total_sales else 0.0, "ratio")


def _write_mall_and_capture(report: Worksheet, metrics: Mapping[str, Any]) -> None:
    mall = _index(metrics["M14"], "dimension", "metrics.M14")
    for dimension, row in MALL_ROWS.items():
        _write(report[f"C{row}"], mall[dimension].get("current"), "amount")
        _write(report[f"D{row}"], mall[dimension].get("wow_change_rate_percent"), "rate_percent")
        _write(report[f"E{row}"], mall[dimension].get("yoy_change_rate_percent"), "rate_percent")
    _write(report["G21"], mall["Total"].get("current"), "amount")
    _write(report["H21"], mall["Total"].get("previous"), "amount")

    capture = _index(metrics["M13"], "dimension", "metrics.M13")
    report["C26"] = "本期"
    report["D26"] = "上期"
    report["E26"] = "环比变化"
    for dimension, row in CAPTURE_ROWS.items():
        _write(report[f"C{row}"], capture[dimension].get("current"), "ratio")
        _write(report[f"D{row}"], capture[dimension].get("base"), "ratio")
        _write(report[f"E{row}"], capture[dimension].get("change_pp"), "pp")
    report["G26"] = "总体本期"
    report["H26"] = "总体上期"
    report["I26"] = "环比变化"
    report["J26"] = "说明"
    _write(report["G27"], capture["Total"].get("current"), "ratio")
    _write(report["H27"], capture["Total"].get("base"), "ratio")
    _write(report["I27"], capture["Total"].get("change_pp"), "pp")
    report["J27"] = "本期较上期"
    for row in range(28, 30):
        for column in "GHIJ":
            report[f"{column}{row}"].value = None


def _write_member_sections(report: Worksheet, metrics: Mapping[str, Any]) -> None:
    points = _index(metrics["M09"], "card_level", "metrics.M09")
    for level, row in POINT_ROWS.items():
        _write(report[f"C{row}"], points.get(level, {}).get("points_balance"), "amount")

    new_members = _index(metrics["M06"], "card_level", "metrics.M06")
    for level, row in NEW_MEMBER_ROWS.items():
        _write(report[f"C{row}"], new_members[level].get("current"), "count")
        _write_change_rate(report[f"D{row}"], new_members[level])
    new_total = new_members.get("Total")
    if new_total is None:
        current_total = sum(float(new_members[level].get("current") or 0) for level in NEW_MEMBER_ROWS)
        base_values = [new_members[level].get("base") for level in NEW_MEMBER_ROWS]
        base_total = None if any(value is None for value in base_values) else sum(float(value) for value in base_values)
        new_total = {"current": current_total, "change_rate_percent": None if base_total in (None, 0) else round((current_total - base_total) / base_total * 100, 1)}
    _write(report["C45"], new_total.get("current"), "count")
    _write_change_rate(report["D45"], new_total)

    stores = _require_sequence(metrics["M07"], "metrics.M07")
    for offset in range(10):
        row = 50 + offset
        if offset >= len(stores):
            for column in "CDEFG":
                report[f"{column}{row}"].value = None
            continue
        item = _require_mapping(stores[offset], "metrics.M07")
        _write(report[f"C{row}"], item.get("store_name"))
        _write(report[f"D{row}"], item.get("current"), "amount")
        _write(report[f"E{row}"], item.get("transactions"), "count")
        _write(report[f"F{row}"], item.get("average_ticket"), "amount")
        _write(report[f"G{row}"], item.get("change_rate_percent"), "rate_percent")
    store_items = [_require_mapping(item, "metrics.M07") for item in stores]
    store_current = sum(float(item.get("current") or 0) for item in store_items)
    store_transactions = sum(int(item.get("transactions") or 0) for item in store_items)
    store_bases = [item.get("base") for item in store_items]
    store_base = None if any(value is None for value in store_bases) else sum(float(value) for value in store_bases)
    _write(report["D60"], store_current, "amount")
    _write(report["E60"], store_transactions, "count")
    _write(report["F60"], None if store_transactions == 0 else store_current / store_transactions, "amount")
    store_rate = None if store_base is None or (store_base == 0 and store_current != 0) else (0.0 if store_base == 0 else round((store_current - store_base) / store_base * 100, 1))
    _write(report["G60"], store_rate, "rate_percent")

    report["C63"] = "会员卡号"
    members = _require_sequence(metrics["M08"], "metrics.M08")
    for offset in range(10):
        row = 64 + offset
        if offset >= len(members):
            for column in "CDEFG":
                report[f"{column}{row}"].value = None
            continue
        item = _require_mapping(members[offset], "metrics.M08")
        _write(report[f"C{row}"], item.get("member_id"))
        report[f"C{row}"].number_format = "@"
        member_type = MEMBER_TYPE_LABELS.get(str(item.get("member_type", "")), str(item.get("member_type", "")))
        _write(report[f"D{row}"], member_type)
        member_type_font = copy(report[f"D{row}"].font)
        member_type_font.color = NEW_MEMBER_FONT_COLOR if member_type == "新会员" else DEFAULT_FONT_COLOR
        report[f"D{row}"].font = member_type_font
        card_level = str(item.get("card_level", ""))
        _write(report[f"E{row}"], card_level)
        card_level_font = copy(report[f"E{row}"].font)
        card_level_font.color = STAR_DIAMOND_FONT_COLOR if card_level == "星钻卡" else DEFAULT_FONT_COLOR
        report[f"E{row}"].font = card_level_font
        stores_text = "、".join(str(value) for value in item.get("stores", []))
        _write(report[f"F{row}"], stores_text)
        _write(report[f"G{row}"], item.get("sales"), "amount")
    _write(report["G74"], sum(float(_require_mapping(item, "metrics.M08").get("sales") or 0) for item in members), "amount")

    active = _index(metrics["M12"], "card_level", "metrics.M12")
    for level, row in ACTIVE_ROWS.items():
        _write(report[f"C{row}"], active[level].get("current"), "count")
        _write(report[f"D{row}"], active[level].get("base"), "count")

    card_sales = _index(metrics["M04"], "card_level", "metrics.M04")
    for level, row in CARD_ROWS.items():
        _write(report[f"C{row}"], card_sales[level].get("current"), "amount")
        _write_change_rate(report[f"D{row}"], card_sales[level])
    card_total = card_sales.get("Total")
    if card_total is None:
        current_total = sum(float(card_sales[level].get("current") or 0) for level in CARD_ROWS)
        base_values = [card_sales[level].get("base") for level in CARD_ROWS]
        base_total = None if any(value is None for value in base_values) else sum(float(value) for value in base_values)
        card_total = {"current": current_total, "change_rate_percent": None if base_total in (None, 0) else round((current_total - base_total) / base_total * 100, 1)}
    _write(report["C99"], card_total.get("current"), "amount")
    _write_change_rate(report["D99"], card_total)

    for merged_range in ("A105:A106", "B105:B107", "B108:B109"):
        if merged_range in {str(item) for item in report.merged_cells.ranges}:
            report.unmerge_cells(merged_range)
    for row in range(106, 110):
        for column in "BCDE":
            _copy_style(report[f"{column}105"], report[f"{column}{row}"])
    rows = _require_sequence(metrics["M05"], "metrics.M05")
    for offset in range(5):
        row = 105 + offset
        if offset >= len(rows):
            for column in "BCDE":
                report[f"{column}{row}"].value = None
            continue
        item = _require_mapping(rows[offset], "metrics.M05")
        _write(report[f"B{row}"], item.get("district"))
        _write(report[f"C{row}"], item.get("card_level"))
        _write(report[f"D{row}"], item.get("current"), "amount")
        _write(report[f"E{row}"], item.get("change_rate_percent"), "rate_percent")


def _apply_report_dates(report: Worksheet, windows: Mapping[str, Any]) -> None:
    current = _window_label(_require_mapping(windows["current"], "windows.current"))
    previous = _window_label(_require_mapping(windows["previous"], "windows.previous"))
    yoy = _window_label(_require_mapping(windows["yoy"], "windows.yoy"))
    report["C2"] = current
    report["D2"] = previous
    report["F2"] = f"同比值 {yoy}"
    for cell in ("C8", "C12", "C16", "C20", "C26"):
        report[cell] = current
    for cell in ("D8", "D12", "D16"):
        report[cell] = previous
    for cell in ("E8", "E12", "E16"):
        report[cell] = "WOW"
    for cell in ("F8", "F12", "F16"):
        report[cell] = "交易笔数"


def _copy_style(source: Any, target: Any) -> None:
    target._style = copy(source._style)
    target.alignment = copy(source.alignment)


def _prepare_new_sheet(sheet: Worksheet, report: Worksheet, headers: list[str], widths: list[float]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 24
    header_source = report["B2"]
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(1, index, header)
        _copy_style(header_source, cell)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = widths[index - 1]
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _style_body_row(sheet: Worksheet, row: int, columns: int, report: Worksheet) -> None:
    body_source = report["B3"]
    for column in range(1, columns + 1):
        cell = sheet.cell(row, column)
        _copy_style(body_source, cell)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _insight_text(insight: Mapping[str, Any]) -> str:
    inference = _replace_brand_names(insight["inference"]).strip()
    recommendation = _require_mapping(insight["recommendation"], "insight.recommendation")
    if recommendation.get("applicable") is True:
        action = _replace_brand_names(recommendation.get("action", "")).strip()
        if action:
            return f"{inference} 建议：{action}"
    else:
        reason = _replace_brand_names(recommendation.get("not_applicable_reason", "")).strip()
        if reason:
            return f"{inference} 建议：不适用（{reason}）"
    return inference


def _business_label(value: Any, labels: Mapping[str, str]) -> str:
    raw = str(value or "").strip()
    return labels.get(raw.upper(), raw)


def _brand_full_name(value: Any) -> str:
    raw = str(value or "").strip()
    return BRAND_FULL_NAMES.get(raw.upper(), raw)


def _replace_brand_names(value: Any) -> str:
    text = str(value or "")
    for source_key, full_name in sorted(BRAND_FULL_NAMES.items(), key=lambda item: len(item[0]), reverse=True):
        text = re.sub(
            rf"(?<![A-Za-z]){re.escape(source_key)}(?![A-Za-z])",
            full_name,
            text,
            flags=re.IGNORECASE,
        )
    return text


def _replace_exact_report_brand_labels(sheet: Worksheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell) or not isinstance(cell.value, str):
                continue
            source_key = cell.value.strip().upper()
            if source_key not in BRAND_FULL_NAMES:
                continue
            cell.value = BRAND_FULL_NAMES[source_key]
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "center"
            cell.alignment = alignment
            sheet.row_dimensions[cell.row].height = max(sheet.row_dimensions[cell.row].height or 15, 30)


def _insight_subject(insight: Mapping[str, Any], groups: Mapping[str, Mapping[str, Any]]) -> str:
    group = groups.get(str(insight.get("anomaly_group_id", "")), {})
    dimension_type = _business_label(group.get("dimension_type"), DIMENSION_LABELS)
    dimension_key = str(group.get("dimension_key", "")).strip()
    if dimension_key.upper() == "TOTAL":
        dimension_key = "总体"
    elif dimension_type == "品牌":
        dimension_key = _brand_full_name(dimension_key)
    metric = _business_label(group.get("metric_id"), METRIC_LABELS)
    scope = dimension_key if dimension_type == "总体" else f"{dimension_type}：{dimension_key}"
    return f"{scope}｜{metric}" if metric else scope


def _write_insight_sheet(
    sheet: Worksheet,
    report: Worksheet,
    insights: list[Mapping[str, Any]] | None,
    groups: list[Mapping[str, Any]],
) -> int:
    headers = ["异常对象/指标", "输出洞察", "置信度", "置信度原因", "验证指标"]
    _prepare_new_sheet(sheet, report, headers, [28, 64, 10, 44, 36])
    if insights is None:
        rows = [("本周异常", "归因节点未返回结果，本次报告仅包含指标与告警。", "不可用", "N04归因节点未成功返回InsightBundle。", "待归因恢复后复核")]
        count = 0
    elif not insights:
        rows = [("本周异常", "本周未生成异常洞察。", "不可用", "上游InsightBundle未包含洞察记录。", "继续监测核心指标")]
        count = 0
    else:
        group_index = {str(group.get("anomaly_group_id", "")): group for group in groups}
        rows = [
            (
                _insight_subject(insight, group_index),
                _insight_text(insight),
                CONFIDENCE_LABELS[str(insight["confidence"]).upper()],
                _replace_brand_names(insight["confidence_reason"]),
                "；".join(_replace_brand_names(item) for item in insight["verification_metrics"]),
            )
            for insight in insights
        ]
        count = len(rows)
    for row_index, values in enumerate(rows, start=2):
        _style_body_row(sheet, row_index, 5, report)
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, value)
        estimated_lines = max(
            (len(str(values[0])) // 18) + 1,
            (len(str(values[1])) // 38) + 1,
            (len(str(values[3])) // 28) + 1,
            (len(str(values[4])) // 24) + 1,
        )
        sheet.row_dimensions[row_index].height = min(180, max(54, 20 * estimated_lines))
    sheet.auto_filter.ref = f"A1:E{max(2, sheet.max_row)}"
    sheet.print_area = f"A1:E{max(2, sheet.max_row)}"
    return count


def _signal_summaries(signals: list[Mapping[str, Any]]) -> dict[str, str]:
    grouped: dict[str, list[str]] = {}
    for signal in signals:
        group_id = str(signal.get("anomaly_group_id", ""))
        signal_type = _business_label(signal.get("signal_type"), SIGNAL_LABELS)
        severity = _business_label(signal.get("severity"), SEVERITY_LABELS)
        direction = _business_label(signal.get("direction"), DIRECTION_LABELS)
        text = f"{signal_type}（{severity}/{direction}）"
        grouped.setdefault(group_id, []).append(text)
    return {group_id: "；".join(items) for group_id, items in grouped.items()}


def _quality_rows(
    metric_bundle: Mapping[str, Any],
    anomaly_bundle: Mapping[str, Any],
    insight_bundle: Mapping[str, Any] | None,
) -> list[tuple[str, str, str, str, str]]:
    rows: list[tuple[str, str, str, str, str]] = []
    for source, records in (
        ("N02", metric_bundle.get("data_quality_status", [])),
        ("N03", anomaly_bundle.get("data_quality_status", [])),
    ):
        for raw in _require_sequence(records, f"{source}.data_quality_status"):
            item = _require_mapping(raw, f"{source}.data_quality_status")
            code = item.get("status_code", item.get("code", ""))
            level = item.get("level", "WARNING")
            scope = item.get("scope", item.get("dimension", ""))
            rows.append((source, str(level), str(scope), str(code), str(item.get("message", ""))))
    if insight_bundle is None:
        rows.append(("N04", "ERROR", "归因洞察", "归因节点不可用", "归因节点未返回结果，本次报告仅保留指标与告警。"))
    else:
        if insight_bundle.get("knowledge_base_status") == "unavailable":
            rows.append(
                (
                    "N04",
                    "WARNING",
                    "活动知识库",
                    "活动知识库不可用",
                    "经营数据归因仍保留；活动相关判断已降级，需后续复核。",
                )
            )
        for raw in _require_sequence(insight_bundle.get("data_quality_status", []), "N04.data_quality_status"):
            item = _require_mapping(raw, "N04.data_quality_status")
            code = item.get("status_code", item.get("code", "归因状态"))
            level = item.get("level", "WARNING")
            scope = item.get("scope", item.get("dimension", "归因洞察"))
            rows.append(("N04", str(level), str(scope), str(code), str(item.get("message", ""))))
    return rows


def _write_alert_sheet(
    sheet: Worksheet,
    report: Worksheet,
    groups: list[Mapping[str, Any]],
    signals: list[Mapping[str, Any]],
    quality: list[tuple[str, str, str, str, str]],
) -> int:
    headers = ["严重度", "维度类型", "维度", "指标", "方向", "当前值", "变化额", "信号摘要"]
    _prepare_new_sheet(sheet, report, headers, [12, 18, 20, 18, 12, 16, 16, 52])
    summaries = _signal_summaries(signals)
    for row_index, group in enumerate(groups, start=2):
        _style_body_row(sheet, row_index, 8, report)
        values = [
            _business_label(group.get("group_severity"), SEVERITY_LABELS),
            _business_label(group.get("dimension_type"), DIMENSION_LABELS),
            "总体" if str(group.get("dimension_key", "")).upper() == "TOTAL" else (
                _brand_full_name(group.get("dimension_key", ""))
                if str(group.get("dimension_type", "")).upper() == "BRAND"
                else group.get("dimension_key", "")
            ),
            _business_label(group.get("metric_id"), METRIC_LABELS),
            _business_label(group.get("direction"), DIRECTION_LABELS),
            group.get("current_value", ""),
            group.get("change_amount", ""),
            summaries.get(str(group.get("anomaly_group_id", "")), ""),
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_index, column, "" if value is None else str(value))
        severity = str(group.get("group_severity", "")).upper()
        if severity == "RED":
            sheet.cell(row_index, 1).fill = PatternFill("solid", fgColor="F4CCCC")
        elif severity == "YELLOW":
            sheet.cell(row_index, 1).fill = PatternFill("solid", fgColor="FFF2CC")
        sheet.row_dimensions[row_index].height = 36

    quality_title_row = max(3, len(groups) + 4)
    sheet.merge_cells(start_row=quality_title_row, start_column=1, end_row=quality_title_row, end_column=8)
    title = sheet.cell(quality_title_row, 1, "数据质量与归因状态")
    _copy_style(report["B31"], title)
    title.font = copy(report["B31"].font)
    title.alignment = Alignment(vertical="center")
    quality_header_row = quality_title_row + 1
    quality_headers = ["来源", "级别", "范围", "状态", "说明"]
    for column, header in enumerate(quality_headers, start=1):
        cell = sheet.cell(quality_header_row, column, header)
        _copy_style(report["B2"], cell)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if not quality:
        quality = [("系统", "INFO", "", "运行正常", "未发现数据质量或归因状态告警")]
    for offset, values in enumerate(quality, start=1):
        row_index = quality_header_row + offset
        _style_body_row(sheet, row_index, 5, report)
        localized = list(values)
        localized[1] = {"INFO": "信息", "WARNING": "警告", "ERROR": "错误"}.get(str(values[1]).upper(), values[1])
        for column, value in enumerate(localized, start=1):
            sheet.cell(row_index, column, value)
        sheet.row_dimensions[row_index].height = 32
    sheet.auto_filter.ref = f"A1:H{max(2, len(groups) + 1)}"
    sheet.print_area = f"A1:H{sheet.max_row}"
    return len(groups)


def _scan_error_tokens(workbook: Any) -> list[dict[str, str]]:
    errors = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                value = cell.value
                if cell.data_type == "f" or (isinstance(value, str) and any(token in value for token in ERROR_TOKENS)):
                    errors.append({"sheet": sheet.title, "cell": cell.coordinate, "value": str(value)})
    return errors


def _scan_missing_report_values(report: Worksheet) -> list[dict[str, str]]:
    missing = []
    for row in report.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell) or not isinstance(cell.value, str):
                continue
            if cell.value.strip().upper() in MISSING_VALUE_TOKENS:
                missing.append({"cell": cell.coordinate, "value": cell.value})
    return missing


def export_weekly_report(
    template_path: str | Path,
    output_path: str | Path,
    metric_bundle: Mapping[str, Any],
    anomaly_bundle: Mapping[str, Any],
    insight_bundle: Mapping[str, Any] | None = None,
) -> ReportExportResult:
    """Copy the fixed weekly template and write report, insight and alert sheets."""
    template = Path(template_path).expanduser().resolve()
    output = Path(output_path).expanduser().resolve()
    if not template.is_file():
        raise ReportExportError("TEMPLATE_NOT_FOUND", "Excel模板不存在", {"path": str(template)})
    if template == output:
        raise ReportExportError("OUTPUT_CONFLICT", "输出路径不能覆盖模板", {"path": str(template)})

    metric_bundle = _require_mapping(metric_bundle, "metric_bundle")
    anomaly_bundle = _require_mapping(anomaly_bundle, "anomaly_bundle")
    metrics, windows = _validate_metric_bundle(metric_bundle)
    groups, signals = _validate_anomaly_bundle(anomaly_bundle)
    insight_payload = None if insight_bundle is None else _require_mapping(insight_bundle, "insight_bundle")
    insights = _validate_insights(insight_payload, groups, signals)

    workbook = load_workbook(template)
    if workbook.sheetnames != [EXPECTED_TEMPLATE_SHEET]:
        raise ReportExportError(
            "TEMPLATE_STRUCTURE_INVALID",
            "模板必须且只能包含Weekly Summary",
            {"sheet_names": workbook.sheetnames},
        )
    report = workbook[EXPECTED_TEMPLATE_SHEET]
    report.title = "报告"
    _clear_formulas_and_invalid_names(workbook, report)
    _apply_report_dates(report, windows)
    _write_district_metrics(report, metrics)
    _write_brand_metrics(report, metrics)
    _write_mall_and_capture(report, metrics)
    _write_member_sections(report, metrics)
    report.sheet_view.showGridLines = False
    report.print_area = "B2:J109"
    report.page_setup.orientation = "landscape"
    report.page_setup.fitToWidth = 1
    report.page_setup.fitToHeight = 0
    report.sheet_properties.pageSetUpPr.fitToPage = True

    insight_sheet = workbook.create_sheet("洞察")
    insight_count = _write_insight_sheet(insight_sheet, report, insights, groups)
    quality = _quality_rows(metric_bundle, anomaly_bundle, insight_payload)
    alert_sheet = workbook.create_sheet("告警")
    alert_count = _write_alert_sheet(alert_sheet, report, groups, signals, quality)

    missing_values = _scan_missing_report_values(report)
    if missing_values:
        raise ReportExportError(
            "REPORT_VALUE_MISSING",
            "报告仍有未计算指标，已阻止输出；请检查积分字段、Mall Sales 品牌名称和比较周期。",
            {"missing_values": missing_values[:30]},
        )
    errors = _scan_error_tokens(workbook)
    if errors:
        raise ReportExportError("WORKBOOK_ERROR_VALUE", "输出工作簿仍包含错误公式或错误值", {"errors": errors[:20]})

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, output)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()

    return ReportExportResult(
        output_path=output,
        week_id=str(metric_bundle.get("week_id", "")),
        sheet_names=tuple(workbook.sheetnames),
        insight_count=insight_count,
        alert_count=alert_count,
        quality_count=len(quality),
    )
