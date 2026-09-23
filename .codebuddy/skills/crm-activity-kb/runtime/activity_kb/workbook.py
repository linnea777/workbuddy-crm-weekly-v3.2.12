from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


DEFAULT_CONFIG = {
    "schema_version": "3.0.0",
    "default_top_k": "5",
    "pre_window_days": "14",
    "post_window_days": "7",
    "weight_time": "0.40",
    "weight_scope": "0.30",
    "weight_activity_type": "0.30",
}


@dataclass(frozen=True)
class WorkbookData:
    path: Path
    activities: list[dict[str, Any]]

    @property
    def config(self) -> dict[str, str]:
        return dict(DEFAULT_CONFIG)


def _normalize_date(value: Any, epoch: datetime) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return from_excel(value, epoch).date().isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"无法识别日期：{value!r}")


def _read_activities(ws) -> list[dict[str, Any]]:
    header_row = None
    for row in range(1, min(ws.max_row, 20) + 1):
        values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        if "activity_id" in values:
            header_row = row
            break
    if header_row is None:
        return []
    headers = [str(ws.cell(header_row, col).value or "").strip() for col in range(1, ws.max_column + 1)]
    rows: list[dict[str, Any]] = []
    for row in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row, col).value for col in range(1, len(headers) + 1)]
        if not any(value not in (None, "") for value in values):
            continue
        record: dict[str, Any] = {}
        for header, value in zip(headers, values):
            if isinstance(value, str):
                value = value.strip()
            if header in {"start_date", "end_date"}:
                value = _normalize_date(value, ws.parent.epoch)
            record[header] = value
        rows.append(record)
    return rows


def load_knowledge_workbook(path: str | Path) -> WorkbookData:
    source = Path(path).expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(f"知识库工作簿不存在：{source}")
    wb = load_workbook(source, data_only=True, read_only=False)
    try:
        if "Activities" not in wb.sheetnames:
            raise ValueError("缺少工作表：Activities")
        return WorkbookData(path=source, activities=_read_activities(wb["Activities"]))
    finally:
        wb.close()
