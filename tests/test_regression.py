import sys
import unittest
import tempfile
import json
from pathlib import Path
from datetime import date,timedelta
from copy import deepcopy
import pandas as pd

ROOT=Path(__file__).resolve().parents[1]
SKILLS=ROOT/'.codebuddy/skills'
for skill in ('crm-metrics','crm-anomaly-detection','crm-data-investigation','crm-report-output'):
    sys.path.insert(0,str(SKILLS/skill/'runtime'))
sys.path.insert(0,str(SKILLS/'crm-weekly-report/scripts'))
from crm_weekly_metrics import CalculationRequest,calculate_metrics
from crm_weekly_metrics.calculator import _prepare,_change
from rich_metrics import build_rich_metrics
from crm_anomaly_detector import detection_input_from_metric_bundle,detect_anomalies,DetectionInput,MetricRecord,HistoryPoint
from governance import enrich_anomaly_bundle
from selection import select_issues,default_plan
from data_investigation.engine import InvestigationEngine
from validate_plan import validate_first_round
from validate_insights import validate_insight_bundle
from build_evidence_ledger import build_ledgers
from issue_contract import validate_tree,validate_disclosure
from common import ContractError

START=date(2026,8,24)


def frames():
    crm=[];mall=[];coverage={'crm':{},'mall':{}}
    for i in list(range(28))+[52]:
        week=START-timedelta(weeks=i)
        coverage['crm'][week.isoformat()]={r:'COMPLETE' for r in ('N','S','W')}
        coverage['mall'][week.isoformat()]={r:'COMPLETE' for r in ('N','S','W')}
        for region in ('N','S','W'):
            brands=['LOUIS VUITTON 路易威登','Tiffany','Burberry','Tiny'] if region=='N' else ['LOUIS VUITTON 路易威登' if region=='S' else 'West']
            for brand in brands:
                amount=10000 if brand=='Tiny' else 200000
                if i==0:
                    if brand=='Tiny':amount=20000
                    elif region=='N':amount=100000
                    elif region=='S':amount=500000
                # current decline staircase for LV
                if brand.startswith('LOUIS') and region=='N' and i in (1,2,3):amount={1:200000,2:300000,3:400000}[i]
                for j in range(5):
                    crm.append(dict(event_time=week,district=region,store_name=brand,member_id=f'{region}-{brand}-{j}',card_level=['银卡','金卡','黑卡','黑钻卡','星钻卡'][j],registration_time=date(2020,1,1),paid_amount=amount/5,points_balance=10))
                mall.append(dict(business_date=week,district=region,store_name=brand,mall_sales=amount*2))
    return pd.DataFrame(crm),pd.DataFrame(mall),coverage


def artifacts(crm=None,mall=None,coverage=None):
    if crm is None:crm,mall,coverage=frames()
    req=CalculationRequest(START,crm,mall,coverage or {})
    prepared=_prepare(req);prepared[0].attrs['coverage']=coverage or {};prepared[1].attrs['coverage']=coverage or {}
    metrics=calculate_metrics(req,prepared=prepared).as_dict();metrics.update(input_signature='synthetic-input',implementation_version='3.2.2')
    rich=build_rich_metrics(prepared[0],prepared[1],metrics,input_signature=metrics['input_signature'])
    anomaly=enrich_anomaly_bundle(detect_anomalies(detection_input_from_metric_bundle(metrics)).to_dict(),metrics)
    return metrics,rich,anomaly


def query(items,region='N',brand='LOUIS VUITTON 路易威登',qid='q',cid='c'):
    return dict(query_id=qid,insight_case_id=cid,anomaly_group_id='g',focus='MEMBER_BEHAVIOR',comparison_basis='WOW',scope=dict(dimension_type='BRAND',dimension_key=brand,region_id=region),analysis_items=items)


class Regression(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.metrics,cls.rich,cls.anomaly=artifacts()

    def test_half_up_boundaries(self):
        for rate in (9.94,9.95,11.95,14.95,29.95):
            for sign in (-1,1):
                expected=sign*float(str(rate))
                from decimal import Decimal,ROUND_HALF_UP
                expected=float(Decimal(str(sign*rate)).quantize(Decimal('.1'),rounding=ROUND_HALF_UP))
                self.assertEqual(_change(100000+sign*rate*1000,100000)['change_rate_percent'],expected)

    def test_card_and_calendar_history(self):
        records=self.metrics['anomaly_records']
        self.assertEqual(len([r for r in records if r['dimension_type']=='MEMBER_TIER' and r['metric_id']=='CRM_SALES']),5)
        self.assertEqual(len([r for r in records if r['dimension_type']=='REGION_MEMBER_TIER']),15)
        for r in records:
            if r['metric_id'] in ('CRM_SALES','CAPTURE_RATIO'):self.assertEqual(len(r['history']),27)

    def test_missing_week_not_zero(self):
        c,m,cov=frames();c=c.loc[c.event_time!=START];cov['crm'].pop(START.isoformat())
        metrics,rich,anomaly=artifacts(c,m,cov)
        self.assertFalse(anomaly['anomaly_groups'])
        self.assertIsNone(rich['crm']['district']['Total']['current'])

    def test_partial_region_and_proven_zero(self):
        c,m,cov=frames();c=c.loc[~((c.event_time==START)&(c.district=='N'))]
        unknown=deepcopy(cov);unknown['crm'][START.isoformat()]['N']='MISSING'
        metrics,_,_=artifacts(c,m,unknown)
        self.assertTrue(all(r['current_value'] is None for r in metrics['anomaly_records'] if r['dimension_type']=='REGION_MEMBER_TIER' and r['dimension_key'].startswith('N|')))
        metrics,_,_=artifacts(c,m,cov)
        self.assertEqual(next(r for r in metrics['anomaly_records'] if r['dimension_type']=='REGION' and r['dimension_key']=='N' and r['metric_id']=='CRM_SALES')['current_value'],0)

    def test_small_red_retained_not_selected(self):
        a=select_issues(deepcopy(self.anomaly),self.metrics,self.rich)
        tiny=next(g for g in a['anomaly_groups'] if g['dimension_key']=='Tiny')
        self.assertEqual(tiny['severity'],'RED');self.assertFalse(tiny['must_investigate'])
        self.assertTrue(a['issues']);self.assertTrue(any(i['issue_type']=='BRAND_TREND' for i in a['issues']))
        plan=default_plan(a,None);self.assertTrue(validate_first_round(a,plan)['ok'])

    def test_band_unique_member_net_and_region(self):
        rich=deepcopy(self.rich)
        rich['crm']['member_week_facts']=[dict(period=period,member_id=member,brand='LOUIS VUITTON 路易威登',region=region,card_tier='BLACK',crm_sales=amount,transactions=1,refund_amount=min(0,amount)) for period,member,region,amount in [('current','a','N',8000),('current','a','N',7000),('current','b','N',10000),('current','refund','N',-3000),('current','wrong','S',999999),('previous','a','N',10000)]]
        result=InvestigationEngine(rich).investigate(query(['CORE_METRICS','AMOUNT_BAND']))['results'][0]
        bands=[f for f in result['facts'] if f['metric_id'] in ('AMOUNT_BAND_CRM_SALES','REFUND_ADJUSTMENT')]
        self.assertEqual(len(bands),7)
        self.assertEqual(sum(f['crm_sales'] for f in bands),22000)
        mid=next(f for f in bands if f['band_lower']==10000)
        self.assertEqual(mid['member_count'],1);self.assertEqual(mid['transaction_count'],2);self.assertEqual(mid['crm_sales'],15000)
        self.assertAlmostEqual(sum(f['sales_share_brand'] or 0 for f in bands),1)
        self.assertEqual(result['analysis_item_statuses'][1]['status'],'COMPLETED')

    def test_ordinary_card_sales_and_no_band_false_success(self):
        e=InvestigationEngine(self.rich)
        r=e.investigate(query(['CORE_METRICS','CARD_TIER'],brand='Tiffany'))['results'][0]
        self.assertEqual(len([f for f in r['facts'] if f['metric_id']=='CARD_TIER_CRM_SALES']),5)
        r=e.investigate(query(['AMOUNT_BAND'],brand='Tiffany'))['results'][0]
        self.assertNotEqual(r['status'],'OK');self.assertEqual(r['analysis_item_statuses'][0]['status'],'UNAVAILABLE')

    def test_fact_reuse_and_cache(self):
        e=InvestigationEngine(self.rich)
        a=e.investigate(query(['CORE_METRICS'],qid='q1',cid='region'))['results'][0]
        b=e.investigate(query(['CORE_METRICS'],qid='q2',cid='trend'))['results'][0]
        self.assertEqual(a['fact_ids'],b['fact_ids']);self.assertEqual(e.cache_hits,1)
        other=deepcopy(self.rich);other['input_signature']='changed'
        d=InvestigationEngine(other).investigate(query(['CORE_METRICS']))['results'][0]
        self.assertNotEqual(a['fact_ids'],d['fact_ids'])

    def test_mixed_signal_direction(self):
        rec=MetricRecord(source_record_id='test',week_id='2026-W35',dimension_type='TOTAL',dimension_key='TOTAL',metric_id='CRM_SALES',current_value=120,previous_value=100,wow_rate=.2,yoy_rate=-.3)
        group=detect_anomalies(DetectionInput(records=(rec,))).to_dict()['anomaly_groups'][0]
        self.assertEqual(group['direction'],'MIXED')

    def test_missing_old_history_does_not_erase_recent_trend(self):
        rec=MetricRecord(source_record_id='test',week_id='2026-W35',dimension_type='TOTAL',dimension_key='TOTAL',metric_id='CRM_SALES',current_value=60,previous_value=80,wow_rate=-.25,history=tuple(HistoryPoint(f'2026-W{w}',v) for w,v in [(30,None),(31,150),(32,120),(33,100),(34,80)]))
        signals=detect_anomalies(DetectionInput(records=(rec,))).to_dict()['anomaly_signals']
        self.assertTrue(any(s['signal_type']=='CONTINUOUS_DECLINE' for s in signals))

    def test_tree_siblings_and_wrong_region(self):
        i=dict(root_node_id='root',comparison_basis='WOW',nodes=[dict(node_id='root',scope={'dimension_type':'REGION','dimension_key':'N'})]+[dict(node_id=b,scope={'dimension_type':'BRAND','dimension_key':b,'region_id':'N'}) for b in ('LV','Tiffany','Burberry')],edges=[dict(parent_node_id='root',child_node_id=b) for b in ('LV','Tiffany','Burberry')])
        validate_tree(i,{},dict(LV=['N'],Tiffany=['N'],Burberry=['N']))
        i['nodes'][1]['scope']['region_id']='S'
        with self.assertRaises(ContractError):validate_tree(i,{},dict(LV=['N'],Tiffany=['N'],Burberry=['N']))

    def test_nine_topics_without_truncation(self):
        insights=[dict(issue_id=f'i{k}',primary_scope={'dimension_type':'BRAND','dimension_key':f'brand{k}'},claims=[dict(claim_id=f'claim{k}',claim_role='TREND',statement='持续下降',primary_insight_id=f'i{k}')]) for k in range(9)]
        bundle=dict(insights=insights,summary_items=[dict(text='必要主题',related_insight_ids=['i0'])],deduplication_review=dict(completed=True,retained_issue_ids=[i['issue_id'] for i in insights],reason='各自不同范围的必要风险'))
        self.assertEqual(validate_disclosure(bundle,{},dict(anomaly_groups=[]))['insight_count'],9)
        bundle.pop('deduplication_review')
        with self.assertRaises(ContractError):validate_disclosure(bundle,{},dict(anomaly_groups=[]))

    def test_cross_region_plan_is_rejected(self):
        a=select_issues(deepcopy(self.anomaly),self.metrics,self.rich)
        plan=default_plan(a,None)
        regional=next(c for c in plan['selected_cases'] if c['issue_type']=='REGION_CONTRIBUTION' and c['primary_scope']['dimension_key']=='N')
        q=next(q for q in regional['queries'] if q['scope'].get('dimension_type')=='BRAND')
        q['scope']['region_id']='S'
        with self.assertRaises(ContractError):validate_first_round(a,plan)

    def test_mall_only_and_brand_capture_do_not_open_business_groups(self):
        for metric in ('MALL_SALES','CAPTURE_RATIO'):
            rec=MetricRecord(source_record_id='test',week_id='2026-W35',dimension_type='BRAND',dimension_key='Tiffany',metric_id=metric,current_value=200,previous_value=100,wow_rate=1)
            self.assertFalse(detect_anomalies(DetectionInput(records=(rec,))).to_dict()['anomaly_groups'])

    def test_partial_mall_compares_common_stores_and_decomposition_reconciles(self):
        c,m,cov=frames()
        m=m.loc[~((m.business_date==START)&(m.district=='N')&(m.store_name=='Tiffany'))]
        # No COMPLETE assertion for this source/week, so missing brand isn't silently zero.
        cov['mall'][START.isoformat()].pop('N')
        metrics,rich,_=artifacts(c,m,cov)
        self.assertIsNotNone(next(r for r in metrics['metrics']['M13'] if r['dimension']=='N')['current'])
        r=InvestigationEngine(rich).investigate(dict(query_id='op',anomaly_group_id='g',scope={'dimension_type':'REGION','dimension_key':'N'},focus='BRAND_OPERATION',analysis_items=['OPERATING_STATUS'],filters={'operating_scope':'COMMON_STORES_TWO_PERIODS'}))['results'][0]
        capture=next(f for f in r['facts'] if f['metric_id']=='CAPTURE_RATIO')
        self.assertEqual(capture['coverage_status'],'MATCHED_SUBSET')
        self.assertIn('mall_component',capture['display_fields'])
        values=rich['matched_operating']['REGION:N']
        c0,c1,m0,m1=[values[k] for k in ('crm_previous','crm_current','mall_previous','mall_current')]
        self.assertAlmostEqual((c1/m1+c0/m0)/2*(m1-m0)+(m1+m0)/2*(c1/m1-c0/m0),c1-c0)

    def test_shared_fact_ledger_has_separate_issue_links(self):
        a=select_issues(deepcopy(self.anomaly),self.metrics,self.rich)
        g=next(g for g in a['anomaly_groups'] if g['dimension_type']=='REGION' and g['dimension_key']=='N' and g['metric_id']=='CRM_SALES')
        plan={k:a[k] for k in ('schema_version','week_id','input_signature','implementation_version')}
        plan['selected_cases']=[dict(insight_case_id=c,anomaly_group_ids=[g['anomaly_group_id']]) for c in ('region-issue','trend-issue')]
        e=InvestigationEngine(self.rich);results=[]
        for cid in ('region-issue','trend-issue'):
            q=query(['CORE_METRICS'],cid=cid,qid=cid);q['anomaly_group_id']=g['anomaly_group_id']
            result=e.investigate(q)['results'][0];result['insight_case_id']=cid;results.append(result)
        payload={k:a[k] for k in ('week_id','input_signature','implementation_version')}
        payload.update(results=results,coverage=[dict(anomaly_group_id=g['anomaly_group_id'],status='INVESTIGATED')])
        ledger=build_ledgers(a,plan,[payload])
        self.assertEqual(ledger['cases'][0]['fact_ids'],ledger['cases'][1]['fact_ids'])
        self.assertEqual(len(ledger['issue_fact_links']),2*ledger['shared_fact_count'])
        poisoned=deepcopy(payload);poisoned['results'][1]['facts'][0]['current']=999
        with self.assertRaises(ContractError):build_ledgers(a,plan,[poisoned])

    def test_nested_contribution_contexts_do_not_collide(self):
        a=select_issues(deepcopy(self.anomaly),self.metrics,self.rich)
        g=next(g for g in a['anomaly_groups'] if g['dimension_type']=='REGION' and g['dimension_key']=='N' and g['metric_id']=='CRM_SALES')
        plan={k:a[k] for k in ('schema_version','week_id','input_signature','implementation_version')}
        plan['selected_cases']=[dict(insight_case_id='nested',anomaly_group_ids=[g['anomaly_group_id']])]
        engine=InvestigationEngine(self.rich);results=[]
        for n,(scope,dimension,top_k) in enumerate([('TOTAL','REGION',3),('REGION:N','BRAND',1),('REGION:N','BRAND',3)]):
            q=dict(query_id=f'nested-{n}',insight_case_id='nested',anomaly_group_id=g['anomaly_group_id'],scope=scope,focus='CONTRIBUTION',comparison_basis='WOW',analysis_items=['CONTRIBUTORS'],next_dimension=dimension,top_k=top_k)
            result=engine.investigate(q)['results'][0];result['insight_case_id']='nested';results.append(result)
        payload={k:a[k] for k in ('week_id','input_signature','implementation_version')}
        payload.update(results=results,coverage=[dict(anomaly_group_id=g['anomaly_group_id'],status='INVESTIGATED')])
        ledger=build_ledgers(a,plan,[payload])
        region_facts=[f for f in ledger['cases'][0]['facts'] if f['scope']=='REGION:N']
        self.assertEqual(len(region_facts),3)
        self.assertEqual(len({f['fact_id'] for f in region_facts}),3)
        self.assertTrue(all(f['values']==region_facts[0]['values'] for f in region_facts))
        # Repeated identical semantics still share their existing IDs.
        again=build_ledgers(a,plan,[dict(payload,results=results+results)])
        self.assertEqual(ledger['shared_fact_count'],again['shared_fact_count'])

    def test_numeric_brand_names_are_not_metric_claims(self):
        from validate_insights import _validate_narrative_numbers, _entity_labels
        source=[dict(scope_details=dict(dimension_type='BRAND',dimension_key='45RPM（NLG-19）'))]
        labels=_entity_labels(source,[])
        _validate_narrative_numbers({'summary':'45RPM（NLG-19）销售-4.50万'}, {'-4.50万'}, 'test', labels)
        with self.assertRaises(ContractError):
            _validate_narrative_numbers({'summary':'45RPM（NLG-19）销售-45.00万'}, {'-4.50万'}, 'test', labels)
        with self.assertRaises(ContractError):
            _validate_narrative_numbers({'summary':'45RPM（NLG-19）销售-4.50万'}, {'-4.50万'}, 'test', set())

    def test_unsupported_continuous_pair_not_reported_as_success(self):
        q=query(['CORE_METRICS']);q['comparison_basis']='CONTINUOUS_PERIOD'
        result=InvestigationEngine(self.rich).investigate(q)['results'][0]
        self.assertEqual(result['status'],'DATA_UNAVAILABLE')
        self.assertEqual(result['analysis_item_statuses'][0]['status'],'UNAVAILABLE')

    def test_numbers_cannot_change_sign_or_borrow_uncited_values(self):
        from validate_insights import _validate_narrative_numbers
        _validate_narrative_numbers({'summary':'下降-10.00万'},{'-10.00万'},'test')
        for text in ('增长+10.00万','下降-99.00万'):
            with self.assertRaises(ContractError):_validate_narrative_numbers({'summary':text},{'-10.00万'},'test')

    def test_quota_omission_rejected_and_two_topics_allowed(self):
        from issue_contract import reject_quota_reason
        with self.assertRaises(ContractError):reject_quota_reason('已经达到八条')
        insights=[dict(issue_id=f'i{k}',primary_scope={'dimension_type':'BRAND','dimension_key':str(k)},claims=[dict(claim_id=f'c{k}',claim_role='TREND',statement='趋势跟进',primary_insight_id=f'i{k}')]) for k in range(2)]
        b=dict(insights=insights,summary_items=[dict(text='趋势待跟进',related_insight_ids=['i0','i1'])])
        self.assertEqual(validate_disclosure(b,{},dict(anomaly_groups=[]))['insight_count'],2)

    def test_sibling_brand_cases_merge_and_export(self):
        from test_pipeline import make_draft
        from run_pipeline import _report_export_projection
        from v3_exporter import export_weekly_report_v3
        from openpyxl import load_workbook
        a=select_issues(deepcopy(self.anomaly),self.metrics,self.rich)
        wanted=[('REGION','N',None),('BRAND','LOUIS VUITTON 路易威登','N'),('BRAND','Tiffany',None),('BRAND','Burberry',None)]
        selected=[]
        for dim,key,reg in wanted:
            selected.append(next(g for g in a['anomaly_groups'] if g['dimension_type']==dim and g['dimension_key']==key and g.get('region_id')==reg and g['metric_id']=='CRM_SALES'))
        ids={g['anomaly_group_id'] for g in selected}
        for g in a['anomaly_groups']:g['material']=g['must_investigate']=g['anomaly_group_id'] in ids
        plan={k:a[k] for k in ('schema_version','week_id','input_signature','implementation_version')}
        plan.update(kb_version=None,selected_cases=[],omitted_anomaly_groups=[{'anomaly_group_id':g['anomaly_group_id'],'reason':'由本次区域主题覆盖'} for g in a['anomaly_groups'] if g['anomaly_group_id'] not in ids])
        results=[];engine=InvestigationEngine(self.rich)
        for i,g in enumerate(selected):
            scope={'dimension_type':g['dimension_type'],'dimension_key':g['dimension_key']}
            if g['dimension_type']=='BRAND':scope['region_id']='N'
            cid=f'branch-{i}';q=dict(query_id=cid,insight_case_id=cid,anomaly_group_id=g['anomaly_group_id'],scope=scope,comparison_basis='WOW',focus='MEMBER_BEHAVIOR',analysis_items=['CORE_METRICS'])
            case=dict(insight_case_id=cid,issue_id=cid,issue_type='REGION_CONTRIBUTION',primary_scope=scope,comparison_basis='WOW',independent_issue_reason='区域主题中的必要分支',selection_reason='重要分支',anomaly_group_ids=[g['anomaly_group_id']],queries=[q],activity_queries=[])
            plan['selected_cases'].append(case)
            r=engine.investigate(q)['results'][0];r['insight_case_id']=cid;results.append(r)
        validate_first_round(a,plan)
        round_result={k:a[k] for k in ('week_id','input_signature','implementation_version')}
        round_result.update(results=results,coverage=[dict(anomaly_group_id=g['anomaly_group_id'],status='INVESTIGATED') for g in selected])
        evidence=build_ledgers(a,plan,[round_result])
        draft=make_draft(plan,evidence);root=draft['insights'][0]
        root.update(supporting_case_ids=[c['insight_case_id'] for c in plan['selected_cases'][1:]],anomaly_group_ids=list(ids),signal_ids=list({sid for g in selected for sid in g['signal_ids']}),fact_ids=list({fid for r in results for fid in r['fact_ids']}),root_node_id='root',nodes=[dict(node_id='root' if i==0 else f'n{i}',scope=c['primary_scope'],comparison_basis='WOW',fact_ids=results[i]['fact_ids']) for i,c in enumerate(plan['selected_cases'])],edges=[dict(parent_node_id='root',child_node_id=f'n{i}') for i in range(1,4)])
        draft['insights']=[root];draft['summary_items']=[dict(text='区域多品牌综合问题',related_insight_ids=[root['issue_id']])]
        self.assertTrue(validate_insight_bundle(a,plan,evidence,draft)['ok'])
        ep,ee,ei=_report_export_projection(plan,evidence,draft)
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'story.xlsx'
            export_weekly_report_v3(SKILLS/'crm-report-output/assets/周报模板.xlsx',path,self.metrics,a,ep,ei,ee)
            self.assertEqual(load_workbook(path)['洞察'].max_row,2)

if __name__=='__main__':unittest.main()
