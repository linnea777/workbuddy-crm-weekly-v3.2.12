"""Real CLI integration with synthetic workbooks; no model or production data."""
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from copy import deepcopy
from openpyxl import load_workbook
from test_regression import ROOT,SKILLS,START,frames,artifacts


def make_draft(plan,evidence):
    draft={k:plan[k] for k in ('schema_version','week_id','input_signature','implementation_version','kb_version')}
    draft.update(insights=[],monitor_items=[],data_quality_items=[],summary_items=[],run_status={'data_quality':'SYNTHETIC_TEST','attribution':'LOW_CONFIDENCE_TEST','knowledge_base':'AVAILABLE' if evidence['knowledge_base_status']!='unavailable' else 'UNAVAILABLE'})
    ledgers={x['insight_case_id']:x for x in evidence['cases']}
    for case in plan['selected_cases']:
        ledger=ledgers[case['insight_case_id']]
        item={k:case[k] for k in ('issue_id','issue_type','primary_scope','comparison_basis','independent_issue_reason','insight_case_id','anomaly_group_ids')}
        item.update(supporting_case_ids=[],summary='销售变化需要结合业务证据跟进',impact_assessment={'type':'UNRESOLVED','statement':'现有聚合只支持结构变化，原因有待核查'},primary_driver={'statement':'结合消费结构核查变化来源','evidence_strength':'UNRESOLVED','evidence_ids':ledger['fact_ids'][:1]},supporting_factors=[],confidence='LOW',confidence_reason='聚合期间匹配，缺少实际经营与活动核查',verification_metrics=['CRM销售额'],recommendation={'applicable':False,'not_applicable_reason':'主因尚不明确'},follow_up_items=['核查该区域本周数据覆盖与经营状态'],signal_ids=ledger['signal_ids'],fact_ids=ledger['fact_ids'],activity_ids=[],claims=[dict(claim_id=case['issue_id']+'-claim',claim_role='TREND' if case['issue_type']=='BRAND_TREND' else 'CONTRIBUTION',statement='本主题需继续跟进',primary_insight_id=case['issue_id'])],related_insight_ids=[])
        item.update(title='销售变化', narrative='本周销售变化需要结合品牌销售与会员消费结构判断。', analysis_note={'basis':'销售及会员消费比较', 'limitations':'缺少经营原因资料'})
        if case['primary_scope']['dimension_type'] in {'TOTAL','REGION'}:
            s=case['primary_scope'];label='TOTAL' if s['dimension_type']=='TOTAL' else 'REGION:'+s['dimension_key']
            op=[f for f in ledger['facts'] if f.get('scope')==label and f.get('metric_id') in {'CRM_SALES','MALL_SALES','CAPTURE_RATIO'} and all(f.get('display_fields',{}).get(k) not in (None,'','N/A') for k in ('current','baseline'))]
            if {f['metric_id'] for f in op} == {'CRM_SALES','MALL_SALES','CAPTURE_RATIO'}:
                statement='Mall销售与积分记录占比共同用于解释销售变化，具体原因仍需核查。'
                item['operating_assessment']={'statement':statement,'fact_ids':[f['fact_id'] for f in op]}
                item['narrative']+=statement
                if s['dimension_type']=='TOTAL':
                    regions={}
                    for region in ('N','S','W'):
                        rf=[f for f in ledger['facts'] if f.get('scope')=='REGION:'+region and f.get('metric_id') in {'CRM_SALES','MALL_SALES','CAPTURE_RATIO'} and all(f.get('display_fields',{}).get(k) not in (None,'','N/A') for k in ('current','baseline'))]
                        if {f['metric_id'] for f in rf} == {'CRM_SALES','MALL_SALES','CAPTURE_RATIO'}:
                            regions[region]={'statement':statement,'fact_ids':[f['fact_id'] for f in rf]}
                        else:
                            reason='该区域同范围经营比较资料不足。';item['narrative']+=reason
                            regions[region]={'unavailable_reason':reason}
                    item['operating_assessment']['regions']=regions
            else:
                reason='同范围Mall与Capture比较资料不足。'
                item['operating_assessment']={'unavailable_reason':reason}
                item['narrative']+=reason
        if case['primary_scope']['dimension_type'] == 'TOTAL':
            names = {'REGION:N':'北区','REGION:S':'南区','REGION:W':'西区'}
            regional = [f for f in ledger['facts'] if f.get('scope') in names and f.get('metric_id') == 'CRM_SALES' and f.get('display_fields', {}).get('impact_amount') is not None]
            item['narrative'] += '；'.join(names[f['scope']] + '销售变化' + f['display_fields']['impact_amount'] for f in regional)
        item['action_kind']='VERIFY'
        item['action_text']='建议团队与门店核实本周数据覆盖和经营状态，再判断是否需要会员到店合作。'
        item['narrative']+='\n\n'+item['action_text']
        draft['insights'].append(item)
    draft.pop('summary_items')
    if len(draft['insights'])>8:draft['deduplication_review']={'completed':True,'retained_issue_ids':[i['issue_id'] for i in draft['insights']],'reason':'构造独立问题用于无截断导出回归'}
    return draft


class Pipeline(unittest.TestCase):
    def test_real_cli_prepare_investigate_finalize_cache_and_nine(self):
        with tempfile.TemporaryDirectory() as td:
            temp=Path(td);data=temp/'data';data.mkdir();run=temp/'run'
            c,m,cov=frames();c.to_excel(data/'crm.xlsx',index=False)
            for region in ('N','S','W'):
                wide=m.loc[m.district.eq(region)].pivot(index='store_name',columns='business_date',values='mall_sales').reset_index()
                wide.rename(columns={'store_name':'Shop' if region=='N' else 'Tenants'},inplace=True)
                wide.insert(0,'No',range(1,len(wide)+1));wide.to_excel(data/f'{region}.xlsx',index=False)
            (data/'coverage.json').write_text(json.dumps(cov))
            script=SKILLS/'crm-weekly-report/scripts/run_pipeline.py'
            def cli(*args):
                p=subprocess.run([sys.executable,str(script),'--output-dir',str(run),*args],text=True,capture_output=True,env=dict(os.environ,AICRM_PYTHON=sys.executable),timeout=180)
                self.assertEqual(p.returncode,0,p.stderr+'\n'+p.stdout)
                return json.loads(p.stdout)
            args=['--phase','prepare','--crm',str(data/'crm.xlsx'),'--mall-n',str(data/'N.xlsx'),'--mall-s',str(data/'S.xlsx'),'--mall-w',str(data/'W.xlsx'),'--week-start',START.isoformat(),'--coverage-json',str(data/'coverage.json')]
            first=cli(*args);second=cli(*args)
            self.assertFalse(first['metrics_cache_hit']);self.assertTrue(second['metrics_cache_hit'])
            activity_plan=json.loads((run/'recommended_plan.json').read_text())
            case=next(c for c in activity_plan['selected_cases'] if c['primary_scope']=={'dimension_type':'REGION','dimension_key':'N'})
            case['queries'].append(dict(query_id='activity-integration',insight_case_id=case['insight_case_id'],anomaly_group_id=case['anomaly_group_ids'][0],focus='MEMBER_BEHAVIOR',comparison_basis='WOW',scope=case['primary_scope'],analysis_items=['ACTIVITY_RESPONSE'],filters={'activity_window':{'start_date':str(START),'end_date':str(START+__import__('datetime').timedelta(days=2))},'baseline_weeks':4}))
            (run/'activity_plan.json').write_text(json.dumps(activity_plan))
            cli('--phase','investigate','--plan',str(run/'activity_plan.json'))
            detailed=json.loads((run/'rich_metrics.json').read_text())
            self.assertTrue(detailed['crm']['member_day_coverage']['available'])
            self.assertTrue(detailed['mall']['day_facts'])
            plan=json.loads((run/'investigation_plan.json').read_text());evidence=json.loads((run/'evidence_ledger.json').read_text())
            self.assertTrue(evidence['cases'])
            synthesis=json.loads((run/'insight_synthesis_context.json').read_text())
            self.assertNotIn('member_id',json.dumps(synthesis));self.assertNotIn('audit_records',json.dumps(synthesis))
            draft=make_draft(plan,evidence)
            (run/'draft.json').write_text(json.dumps(draft,ensure_ascii=False))
            final=cli('--phase','finalize','--insights',str(run/'draft.json'))
            book=load_workbook(final['report']);self.assertEqual(book.sheetnames,['报告','洞察','分析说明'])
            self.assertEqual(book['洞察'].max_row,len(draft['insights'])+1)
            self.assertEqual(book['洞察'].max_column,2)
            self.assertEqual(book['洞察']['B2'].value,draft['insights'][0]['narrative'])
            # Brand collaborations and member experiences must survive the real
            # validation/export path with the same evidence requirements.
            for channel in ('会员体验活动', '品牌联合活动'):
                with self.subTest(channel=channel):
                    event_draft=deepcopy(draft)
                    item=event_draft['insights'][0]
                    item['confidence']='MEDIUM'
                    item['primary_driver']['evidence_strength']='AGGREGATE'
                    item['recommendation']={'applicable':True,'target':'相关门店','action':'结合品牌上新讨论会员到店体验','channel':channel,'timing':'门店确认上新档期后','fact_ids':item['fact_ids'][:1]}
                    item['action_kind']='CONDITIONAL'
                    item['action_condition']='门店确认上新档期后'
                    item['action_text']='门店确认上新档期后，建议团队结合品牌上新讨论会员到店体验。'
                    item['recommendation']['action']=item['action_text']
                    item['narrative']=item['narrative'].rsplit('\n\n',1)[0]+'\n\n'+item['action_text']
                    (run/'event_draft.json').write_text(json.dumps(event_draft,ensure_ascii=False))
                    result=cli('--phase','finalize','--insights',str(run/'event_draft.json'))
                    self.assertEqual(load_workbook(result['report'])['洞察']['B2'].value,item['narrative'])
            # Build nine distinct scope cases, all using actual synthetic aggregates.
            anomaly=json.loads((run/'anomaly_bundle.json').read_text())
            candidates=[g for g in anomaly['anomaly_groups'] if g['metric_id']=='CRM_SALES' and g['dimension_type'] in ('REGION','MEMBER_TIER','REGION_MEMBER_TIER')][:9]
            self.assertEqual(len(candidates),9)
            chosen={g['anomaly_group_id'] for g in candidates}
            for g in anomaly['anomaly_groups']:
                g['must_investigate']=g['material']=g['anomaly_group_id'] in chosen
            (run/'anomaly_bundle.json').write_text(json.dumps(anomaly))
            nine_plan={k:plan[k] for k in ('schema_version','week_id','input_signature','implementation_version','kb_version')}
            nine_plan.update(selected_cases=[],omitted_anomaly_groups=[{'anomaly_group_id':g['anomaly_group_id'],'reason':'构造场景中由指定主题覆盖'} for g in anomaly['anomaly_groups'] if g['anomaly_group_id'] not in chosen])
            for n,g in enumerate(candidates):
                cid=f'INDEPENDENT-{n}';scope={'dimension_type':g['dimension_type'],'dimension_key':g['dimension_key']}
                q=dict(query_id=cid+'-query',insight_case_id=cid,anomaly_group_id=g['anomaly_group_id'],focus='MEMBER_BEHAVIOR',comparison_basis='WOW',scope=scope,analysis_items=['CORE_METRICS'])
                nine_plan['selected_cases'].append(dict(insight_case_id=cid,issue_id=cid,issue_type='INDEPENDENT_RISK',primary_scope=scope,comparison_basis='WOW',independent_issue_reason='不同范围的独立跟进问题：构造验收场景',selection_reason='构造独立问题',anomaly_group_ids=[g['anomaly_group_id']],queries=[q],activity_queries=[]))
            (run/'nine_plan.json').write_text(json.dumps(nine_plan))
            cli('--phase','investigate','--plan',str(run/'nine_plan.json'))
            evidence=json.loads((run/'evidence_ledger.json').read_text())
            many=make_draft(nine_plan,evidence)
            (run/'many.json').write_text(json.dumps(many,ensure_ascii=False))
            final=cli('--phase','finalize','--insights',str(run/'many.json'))
            self.assertEqual(load_workbook(final['report'])['洞察'].max_row,10)

if __name__=='__main__':unittest.main()
